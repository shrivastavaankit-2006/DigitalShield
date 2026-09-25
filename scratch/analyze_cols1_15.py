import openpyxl
from collections import Counter

excel_path = r"c:\Users\ASHISH SHRIVASTAV\MyCEPProject\src\data\cep.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Sheet1']

headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
rows = []
for row in range(2, sheet.max_row + 1):
    vals = [sheet.cell(row=row, column=col).value for col in range(1, sheet.max_column + 1)]
    if any(v is not None and str(v).strip() != "" for v in vals):
        rows.append(vals)

total_respondents = len(rows)

for col_idx in range(0, 15):
    header = headers[col_idx]
    vals = [r[col_idx] for r in rows]
    clean_vals = [str(v).strip() for v in vals if v is not None and str(v).strip() != ""]
    print(f"\n==========================================")
    print(f"Col {col_idx+1}: {header}")
    
    # Check multi-select
    # Col 2 is Name - do not print individual names for privacy
    if col_idx == 1:
        print(f"Name column: {len(clean_vals)} names present. (PII - Will be excluded)")
        continue
        
    raw_counts = Counter(clean_vals)
    for k, v in raw_counts.most_common():
        print(f"  {k}: {v} ({(v/total_respondents)*100:.1f}%)")
        
    # If it's a multi-select column like Col 5 (platforms), Col 8 (doubtful info), Col 9 (difficult to identify), Col 10 (makes you believe), Col 12 (warning signs)
    if col_idx in [4, 7, 8, 9, 11]:
        split_items = []
        for val in clean_vals:
            # Note: "All of the above" in Col 12 represents all warning signs
            parts = [p.strip() for p in val.split(',') if p.strip()]
            split_items.extend(parts)
        split_counts = Counter(split_items)
        print("  --> Split counts:")
        for k, v in split_counts.most_common():
            print(f"      {k}: {v} respondents ({(v/total_respondents)*100:.1f}%)")
