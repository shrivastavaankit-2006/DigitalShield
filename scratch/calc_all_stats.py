import openpyxl
from collections import Counter
import json

excel_path = r"c:\Users\ASHISH SHRIVASTAV\MyCEPProject\src\data\cep.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Sheet1']

rows = []
for r in range(2, sheet.max_row + 1):
    vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
    if any(v is not None and str(v).strip() != "" for v in vals):
        rows.append(vals)

N = len(rows)
print(f"Total N = {N}")

def pct(count, total=N):
    return round((count / total) * 100, 1)

# Col 1: Age
age_counts = Counter([str(r[0]).strip() for r in rows if r[0] is not None])
print("Age:", age_counts)

# Col 3: Occupation
occ_counts = Counter([str(r[2]).strip() for r in rows if r[2] is not None])
print("Occupation:", occ_counts)

# Col 4: Internet frequency
net_counts = Counter([str(r[3]).strip() for r in rows if r[3] is not None])
print("Net:", net_counts)

# Col 5: Platforms (split)
platforms = []
for r in rows:
    if r[4]:
        for p in str(r[4]).split(','):
            platforms.append(p.strip())
plat_counts = Counter(platforms)
print("Platforms:", plat_counts)

# Col 6: False news encountered
fake_counts = Counter([str(r[5]).strip() for r in rows if r[5] is not None])
print("Fake news:", fake_counts)

# Col 7: Verification frequency
verif_counts = Counter([str(r[6]).strip() for r in rows if r[6] is not None])
print("Verification:", verif_counts)

# Col 8: Action on doubtful info
doubt_actions = []
for r in rows:
    if r[7]:
        for d in str(r[7]).split(','):
            doubt_actions.append(d.strip())
doubt_counts = Counter(doubt_actions)
print("Doubt actions:", doubt_counts)

# Col 9: Difficult to identify as fake
diff_items = []
for r in rows:
    if r[8]:
        for d in str(r[8]).split(','):
            diff_items.append(d.strip())
diff_counts = Counter(diff_items)
print("Diff to identify:", diff_counts)

# Col 10: Makes believe
bel_items = []
for r in rows:
    if r[9]:
        for b in str(r[9]).split(','):
            bel_items.append(b.strip())
bel_counts = Counter(bel_items)
print("Belief factors:", bel_counts)

# Col 11: Suspicious message received
susp_counts = Counter([str(r[10]).strip() for r in rows if r[10] is not None])
print("Suspicious msg:", susp_counts)

# Col 12: Warning signs
warn_items = []
for r in rows:
    if r[11]:
        for w in str(r[11]).split(','):
            warn_items.append(w.strip())
warn_counts = Counter(warn_items)
print("Warning signs:", warn_counts)

# Col 13: OTP share
otp_counts = Counter([str(r[12]).strip() for r in rows if r[12] is not None])
print("OTP share:", otp_counts)

# Col 14: Before click unknown link
link_counts = Counter([str(r[13]).strip() for r in rows if r[13] is not None])
print("Before click link:", link_counts)

# Col 15: Fraud experience
fraud_counts = Counter([str(r[14]).strip() for r in rows if r[14] is not None])
print("Fraud exp:", fraud_counts)

# Col 16: Confidence (1-5)
conf_counts = Counter([int(r[15]) for r in rows if r[15] is not None])
print("Confidence:", conf_counts)

# Col 17: Importance of awareness (1-5)
imp_counts = Counter([int(r[16]) for r in rows if r[16] is not None])
print("Importance:", imp_counts)

# Col 18: Topics want to learn
topics = []
for r in rows:
    if r[17]:
        for t in str(r[17]).split(','):
            topics.append(t.strip())
topic_counts = Counter(topics)
print("Topics to learn:", topic_counts)

# Col 19: Awareness helps
help_counts = Counter([str(r[18]).strip() for r in rows if r[18] is not None])
print("Awareness helps:", help_counts)

# Col 20: Preferred method
meth_counts = Counter([str(r[19]).strip() for r in rows if r[19] is not None])
print("Preferred method:", meth_counts)
