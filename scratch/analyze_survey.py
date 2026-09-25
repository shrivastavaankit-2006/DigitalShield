import openpyxl
from collections import Counter
import json

excel_path = r"c:\Users\ASHISH SHRIVASTAV\MyCEPProject\src\data\cep.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['Sheet1']

headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
rows = []
for row in range(2, sheet.max_row + 1):
    vals = [sheet.cell(row=row, column=col).value for col in range(1, sheet.max_column + 1)]
    if any(v is not None and str(v).strip() != "" for v in vals):
        rows.append(vals)

total_respondents = len(rows)
print(f"Total valid respondents: {total_respondents}")

analysis = {}

for col_idx, header in enumerate(headers):
    vals = [r[col_idx] for r in rows]
    clean_vals = [str(v).strip() for v in vals if v is not None and str(v).strip() != ""]
    
    print(f"\n==========================================")
    print(f"Col {col_idx+1}: {header}")
    print(f"Non-empty count: {len(clean_vals)} / {total_respondents}")
    
    # Check if multiple choice (contains commas)
    has_comma = any(',' in v for v in clean_vals)
    
    # Raw counter
    raw_counts = Counter(clean_vals)
    print("Raw frequencies:")
    for k, v in raw_counts.most_common(10):
        print(f"  {k}: {v} ({(v/total_respondents)*100:.1f}%)")
    if len(raw_counts) > 10:
        print(f"  ... and {len(raw_counts) - 10} more unique values")
        
    if has_comma:
        split_items = []
        for val in clean_vals:
            # Note: split by comma, strip whitespace
            parts = [p.strip() for p in val.split(',') if p.strip()]
            split_items.extend(parts)
        split_counts = Counter(split_items)
        print("\nSplit item frequencies (Multi-select):")
        for k, v in split_counts.most_common():
            print(f"  {k}: {v} respondents ({(v/total_respondents)*100:.1f}%)")
