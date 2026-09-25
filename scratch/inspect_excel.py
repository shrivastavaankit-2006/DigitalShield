import openpyxl
import json
import os

excel_path = r"c:\Users\ASHISH SHRIVASTAV\MyCEPProject\src\data\cep.xlsx"

wb = openpyxl.load_workbook(excel_path, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    print(f"\n--- Sheet: {sheetname} ---")
    print(f"Max rows: {sheet.max_row}, Max cols: {sheet.max_column}")
    
    # Read headers
    headers = []
    for col in range(1, sheet.max_column + 1):
        headers.append(sheet.cell(row=1, column=col).value)
    print("\nHeaders:")
    for idx, h in enumerate(headers, 1):
        print(f"  Col {idx}: {repr(h)}")
        
    # Count rows with data
    data_rows = []
    for row in range(2, sheet.max_row + 1):
        row_vals = [sheet.cell(row=row, column=col).value for col in range(1, sheet.max_column + 1)]
        if any(v is not None and str(v).strip() != "" for v in row_vals):
            data_rows.append(row_vals)
    print(f"\nTotal non-empty data rows: {len(data_rows)}")
    
    if len(data_rows) > 0:
        print("\nFirst row sample:")
        for idx, (h, v) in enumerate(zip(headers, data_rows[0]), 1):
            print(f"  {idx}. {h} -> {repr(v)}")
            
        print("\nLast row sample:")
        for idx, (h, v) in enumerate(zip(headers, data_rows[-1]), 1):
            print(f"  {idx}. {h} -> {repr(v)}")
